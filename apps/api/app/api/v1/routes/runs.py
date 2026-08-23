"""Starting and reading runs - org-scoped, persisted.

Every run dials for real. There is no dry-run mode (CLAUDE.md, ADR-3).

**The only per-dial guard is the suppression list.** The per-run ceiling, the
allowlist, per-organisation rate limiting, the daily budget and the
per-teammate call-count allocation were all removed from the dial gate at the
product owner's direction while a replacement security layer is designed - see
`domain/safety.py`'s module docstring. `RunDialer` throttles nothing beyond
that.

**Usage credit is still checked, once, before a run starts** (not per dial):
the organisation-wide balance and the acting teammate's own share of it, both
refusing before a single contact is dialled if either is exhausted. This is a
different system from the removed cost guards above - money, not a call
count - and was not part of that removal.
"""

from __future__ import annotations

import logging
import uuid
from collections.abc import Iterable
from typing import Annotated, Any
from uuid import UUID

from fastapi import (
    APIRouter,
    BackgroundTasks,
    Depends,
    File,
    HTTPException,
    UploadFile,
    status,
)
from pydantic import BaseModel, Field

from app.auth.dependencies import CurrentUser, RequirePermission, current_user
from app.auth.permissions import Permission
from app.core.config import config
from app.core.logging import CallContext
from app.database import database
from app.database.repositories import credits as credits_repo
from app.database.repositories import escalations as escalations_repo
from app.database.repositories import run_numbers as run_numbers_repo
from app.database.repositories import runs as runs_repo
from app.database.repositories import suppressions as suppressions_repo
from app.database.repositories import voice_agents as voice_agents_repo
from app.domain.entities import NEEDS_A_PERSON_DISPOSITIONS, CallOutcome, Contact
from app.domain.entitlements import (
    AgentStanding,
    check_agent_usable,
    check_credit_available,
    check_member_credit_cap,
    usable_agent_ids,
)
from app.domain.run_state import RunStatus, is_stopping
from app.domain.safety import mask, phone_hash
from app.domain.spreadsheet import (
    MAX_ROWS,
    SheetTooLarge,
    parse_rows,
    to_contact_payload,
)
from app.services import billing
from app.services import credit as credit_service
from app.services.run_control import RunNotStoppable, StopSignal, request_stop
from app.services.run_dialer import RunDialer
from app.services.run_dispatch import RunNotDispatchable, RunPlan, resolve_run_plan

log = logging.getLogger("app.api.v1.runs")

router = APIRouter(prefix="/api/v1/runs", tags=["runs"])


class ContactIn(BaseModel):
    name: str
    phone: str
    region: str | None = None
    language: str | None = None
    context: dict[str, Any] = Field(default_factory=dict)


class RunRequest(BaseModel):
    """What starting a run needs.

    The number lives here rather than on the agent (ADR-8), which is what lets
    one agent dial through any carrier the organisation has connected.
    """

    voice_agent_id: UUID
    #: One, several, or every verified number. A run with none is refused.
    number_ids: list[UUID] = Field(default_factory=list)
    contacts: list[ContactIn]
    #: A human label for the run list. Optional - runs are also identifiable by
    #: their agent and start time.
    name: str | None = Field(default=None, max_length=120)
    #: Appended to every prompt in this run, so a one-off instruction does not
    #: mean editing an agent the whole organisation shares.
    run_instruction: str | None = Field(default=None, max_length=2000)
    allocation_strategy: str = "round_robin"


def _deduplicate(contacts: Iterable[Contact]) -> list[Contact]:
    """Collapse contacts that would share one `call_outcomes` row.

    That table is keyed on (run_id, contact_name, phone_masked), so a CSV
    listing the same person twice produces one row however many times it is
    dialled - while `runs.total` counted both. The settled count could then
    never reach the total and the run stayed "running" forever.

    Deduplicating beats loosening the key: a run should not dial the same
    person twice anyway, and the row that would have been overwritten was a
    real call whose transcript was being discarded.
    """
    seen: set[tuple[str, str]] = set()
    unique: list[Contact] = []
    for contact in contacts:
        key = (contact.name, contact.phone)
        if key in seen:
            log.info("skipping a repeated contact in this run's list")
            continue
        seen.add(key)
        unique.append(contact)
    return unique


async def _run_and_persist(
    *,
    run_id: str,
    org_id: UUID,
    auth_user_id: str,
    plan: RunPlan,
    contacts: list[Contact],
    suppressed_hashes: frozenset[str],
) -> None:
    dialer = RunDialer(
        suppressed_hashes=suppressed_hashes,
        run_id=run_id,
        # Read from `runs.stop_requested_at`, not from process memory: the stop
        # request lands on whichever worker serves it, which is rarely this one
        # (`services/run_control.py`).
        should_stop=StopSignal(run_id, auth_user_id).is_stopped,
        # The two things nothing ever passed before, which is why every dial was
        # refused before the phone rang.
        lines=plan.lines,
        voice_agent=plan.voice_agent,
        allocation_strategy=plan.allocation_strategy,
        run_instruction=plan.run_instruction,
    )

    async def on_progress(outcome: CallOutcome) -> None:
        record = outcome.model_dump(mode="json")
        record["provider_call_id"] = record.pop("run_id", None)
        async with database.as_user(auth_user_id) as conn:
            call_outcome_id = await runs_repo.append_outcome(
                conn, run_id=run_id, org_id=org_id, outcome=record
            )
            if outcome.disposition in NEEDS_A_PERSON_DISPOSITIONS:
                await escalations_repo.create_for_outcome(
                    conn, org_id=org_id, run_id=run_id, call_outcome_id=call_outcome_id
                )

    # Every log line this run produces - in this module and `RunDialer` -
    # carries `run_id`/`org_id` for the run's whole lifetime, so one run's
    # lines can be grepped together regardless of which contact or module
    # emitted them.
    with CallContext(run_id=run_id, org_id=str(org_id)):
        try:
            await dialer.run(plan.agent, contacts, on_progress=on_progress)
            # Deliberately not `finish_run()`. Origination returns when a call
            # is answered, not when it ends, so at this point conversations are
            # still running - closing the run here would show it completed
            # while its own rows still read "In conversation…". Whichever
            # worker callback settles the last contact closes it instead
            # (`finish_if_all_settled`). This still handles the case where every
            # contact was blocked or failed to dial: those settle immediately,
            # so the check below closes the run right away.
            async with database.as_user(auth_user_id) as conn:
                await runs_repo.finish_if_all_settled(
                    conn, run_id, stale_after_seconds=int(config.poll_timeout_seconds) * 2
                )
        except Exception as exc:
            log.exception("run %s failed", run_id)
            async with database.as_user(auth_user_id) as conn:
                await runs_repo.finish_run(conn, run_id, error=f"{type(exc).__name__}: {exc}")


@router.post("", status_code=status.HTTP_200_OK)
async def start_run(
    req: RunRequest,
    background: BackgroundTasks,
    user: Annotated[CurrentUser, Depends(RequirePermission(Permission.RUNS_START))],
) -> dict[str, Any]:
    # Resolved before anything is written: a run that cannot dial must not
    # exist as a row that never starts, and every refusal here names the thing to
    # go and fix.
    async with database.as_user(user.auth_user_id) as conn:
        try:
            plan = await resolve_run_plan(
                conn,
                org_id=user.org_id,
                voice_agent_id=req.voice_agent_id,
                number_ids=req.number_ids,
                allocation_strategy=req.allocation_strategy,
                run_instruction=req.run_instruction,
            )
        except RunNotDispatchable as exc:
            raise HTTPException(status_code=400, detail=str(exc)) from exc

    if not req.contacts:
        raise HTTPException(status_code=400, detail="At least one contact is required.")

    try:
        contacts = _deduplicate(Contact(**c.model_dump()) for c in req.contacts)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    run_id = uuid.uuid4().hex[:12]
    async with database.as_user(user.auth_user_id) as conn:
        # Usage credit, checked once before the run starts rather than per dial.
        #
        # **The limit this accepts, stated rather than hidden:** a run that starts
        # with credit can overspend it, because the money only moves when each
        # call settles and nothing re-checks the balance mid-run. The exposure is
        # bounded by how many contacts are in the run times the rate, and a call
        # already talking to a person is never cut off for money anyway, so a
        # mid-run balance check could only refuse calls that had not started.
        #
        # A per-dial persisted hold would close that gap and needs a database
        # connection inside `RunDialer`, which is deliberately I/O-free. That is
        # a real change to its shape, not a line here.
        effective_plan = await billing.resolve_plan(conn, user.org_id, user.org_plan_id)

        # The plan's agent limit, enforced where it actually binds. `usable_agent_ids`
        # has always decided which agents the Agents tab shows as locked, but nothing
        # checked it here - so a downgrade greyed an agent out in one screen and the
        # run composer would still dial with it. That made the limit a one-time toll
        # rather than an entitlement: subscribe, build ten agents, downgrade, keep
        # running all ten (`check_agent_usable`'s own docstring says so, and had no
        # caller until now).
        agent_rows = await voice_agents_repo.list_org_agents(conn, user.org_id)
        billing.refuse(
            check_agent_usable(
                agent_id=str(req.voice_agent_id),
                usable=usable_agent_ids(
                    [
                        AgentStanding(
                            agent_id=str(r["id"]),
                            created_at=r["created_at"],
                            kept_at=r["kept_at"],
                        )
                        for r in agent_rows
                    ],
                    limit=effective_plan.entitlements.max_voice_agents,
                ),
                plan_name=billing.plan_name(effective_plan.plan_id),
                limit=effective_plan.entitlements.max_voice_agents,
            )
        )

        if effective_plan.entitlements.monthly_credit_paise is not None:
            # An unsubscribed plan is granted lazily, because nothing else will:
            # `grant_for_period` runs off subscription webhooks and Free has no
            # subscription. Idempotent per calendar month, so this is a no-op after
            # the first run of the month.
            await credit_service.ensure_period_credit(
                conn,
                org_id=user.org_id,
                entitlements=effective_plan.entitlements,
                has_subscription=billing.subscription_grants_credit(effective_plan),
            )
            balance = await credits_repo.balance(conn, user.org_id)
            billing.refuse(
                check_credit_available(
                    balance_paise=balance,
                    plan_name=billing.plan_name(effective_plan.plan_id),
                )
            )

            # A ceiling on this teammate's own share of the pool above,
            # independent of whether the org-wide balance would allow the run -
            # both are checked, and either can refuse. `cap_paise=None` means
            # nobody has set a share for this person, so only the org-wide
            # balance just checked governs them.
            cap_paise, spent_paise = await credit_service.member_credit_cap_status(
                conn, org_id=user.org_id, user_id=user.id
            )
            billing.refuse(
                check_member_credit_cap(
                    cap_paise=cap_paise,
                    spent_paise=spent_paise,
                    plan_name=billing.plan_name(effective_plan.plan_id),
                )
            )

        # Resolved once here, not once per contact - the one guard that still
        # stands before a dial.
        suppressed: set[str] = set()
        for contact in contacts:
            digest = phone_hash(contact.phone)
            if await suppressions_repo.is_suppressed(conn, user.org_id, digest):
                suppressed.add(digest)

        await runs_repo.create_run(
            conn,
            run_id=run_id,
            org_id=user.org_id,
            voice_agent_id=req.voice_agent_id,
            total=len(contacts),
            started_by=user.id,
            name=req.name,
            run_instruction=req.run_instruction,
            allocation_strategy=req.allocation_strategy,
        )
        # Which lines this run may dial from, recorded before it starts: the run
        # is the permanent record of where its calls came from, and resolving
        # that later from a number that has since been retired would lose it.
        await run_numbers_repo.attach(
            conn,
            run_id=run_id,
            org_id=user.org_id,
            number_ids=[line.number_id for line in plan.lines],
        )

    background.add_task(
        _run_and_persist,
        run_id=run_id,
        org_id=user.org_id,
        auth_user_id=user.auth_user_id,
        plan=plan,
        contacts=contacts,
        suppressed_hashes=frozenset(suppressed),
    )
    return {"run_id": run_id, "total": len(contacts)}


@router.get("")
async def list_runs(user: Annotated[CurrentUser, Depends(current_user)]) -> list[dict[str, Any]]:
    async with database.as_user(user.auth_user_id) as conn:
        rows = await runs_repo.list_runs(conn, user.org_id)
    return [
        {
            "id": r["id"],
            "voice_agent_id": str(r["voice_agent_id"]) if r["voice_agent_id"] else None,
            "agent_name": r["agent_name"],
            "name": r["name"],
            "total": r["total"],
            "status": r["status"],
            "started_at": r["started_at"].isoformat(),
            "finished_at": r["finished_at"].isoformat() if r["finished_at"] else None,
            "error": r["error"],
            "completed": r["completed"],
            # Derived once here rather than in each of the three surfaces that
            # render a run's state, so the list, the detail page and the
            # dashboard cannot disagree about what "Stopping…" means.
            "stopping": is_stopping(RunStatus(r["status"]), r["stop_requested_at"]),
            "started_by": str(r["started_by"]) if r["started_by"] else None,
            "started_by_name": r["started_by_name"],
            "started_by_avatar_url": r["started_by_avatar_url"],
        }
        for r in rows
    ]


class TeamMemberSummary(BaseModel):
    user_id: str | None
    name: str | None
    avatar_url: str | None
    total_runs: int
    total_calls: int


@router.get(
    "/team-summary",
    response_model=list[TeamMemberSummary],
    dependencies=[Depends(RequirePermission(Permission.RUNS_READ_TEAM))],
)
async def team_summary(user: Annotated[CurrentUser, Depends(current_user)]) -> list[TeamMemberSummary]:
    """Call volume per teammate, for the admin/owner/viewer dashboard chart."""
    async with database.as_user(user.auth_user_id) as conn:
        rows = await runs_repo.summarize_by_member(conn, user.org_id)
    return [
        TeamMemberSummary(
            user_id=str(r["started_by"]) if r["started_by"] else None,
            name=r["started_by_name"],
            avatar_url=r["started_by_avatar_url"],
            total_runs=r["total_runs"],
            total_calls=r["total_calls"],
        )
        for r in rows
    ]


@router.get("/{run_id}")
async def get_run(
    run_id: str, user: Annotated[CurrentUser, Depends(current_user)]
) -> dict[str, Any]:
    async with database.as_user(user.auth_user_id) as conn:
        run = await runs_repo.get_run(conn, user.org_id, run_id)
        if run is None:
            raise HTTPException(status_code=404, detail="Run not found")
        outcome_rows = await runs_repo.list_outcomes(conn, run_id)
        number_rows = await run_numbers_repo.list_for_run(
            conn, run_id=run_id, org_id=user.org_id
        )

    numbers = [
        {
            "id": str(n["number_id"]),
            # The join is a left join, so `phone_e164` is null if the number row
            # ever goes missing. It cannot be deleted today (no delete grant),
            # but a run is permanent and this response must not 500 on a future
            # schema where it can.
            "phone_masked": mask(n["phone_e164"]) if n["phone_e164"] else "unknown",
            "provider": n["provider"],
            "label": n["label"],
            # Whether the line is still usable. A run from last month may well
            # have been placed from a number since retired, and saying so is
            # more useful than showing it as though it were still live.
            "status": n["status"],
        }
        for n in number_rows
    ]

    outcomes = [
        {
            "contact_name": o["contact_name"],
            "phone_masked": o["phone_masked"],
            "voice_agent_id": str(run["voice_agent_id"]) if run["voice_agent_id"] else None,
            "agent_name": run["agent_name"],
            "status": o["status"],
            "run_id": run_id,
            "provider_call_id": o["provider_call_id"],
            "transcript": o["transcript"],
            "summary": o["summary"],
            "sentiment": o["sentiment"],
            "sentiment_reason": o["sentiment_reason"],
            "extracted": o["extracted"],
            "disposition": o["disposition"],
            "disposition_reason": o["disposition_reason"],
            "error": o["error"],
            "duration_seconds": o["duration_seconds"],
            "created_at": o["created_at"].isoformat(),
            "task_completed": o["task_completed"],
            "completion_confidence_score": o["completion_confidence_score"],
            "completion_confidence_label": o["completion_confidence_label"],
            "evidence": o["evidence"],
            "attempts": o["attempts"],
            # The fields the agent was actually sent to collect. Persisted
            # since `ISSUES.md` #155 - the column, the domain model, the API
            # type and `TranscriptView` all existed, and only the write and
            # this serialisation were missing, so every run showed an empty
            # "What we asked for".
            "collected": o["collected"],
            "missing_required_fields": o["missing_required_fields"],
            "handoff_questions": o["handoff_questions"],
            "from_number_masked": o["from_number_masked"],
        }
        for o in outcome_rows
    ]
    resolved = [o for o in outcomes if o["disposition"] != "in_flight"]
    escalated = sum(1 for o in resolved if o["disposition"] == "escalated")

    started_by = run["started_by"]
    stopped_at = run["stop_requested_at"]
    return {
        "id": run["id"],
        "voice_agent_id": str(run["voice_agent_id"]) if run["voice_agent_id"] else None,
        "agent_name": run["agent_name"],
        # All three were accepted, persisted and used by the dial path, and none
        # of them were ever read back - so a finished run could not be audited
        # against what it was actually told to do. `run_instruction` matters
        # most: it is appended to every prompt in the run, and "why did this
        # batch say that?" was unanswerable without it.
        "name": run["name"],
        "run_instruction": run["run_instruction"],
        "allocation_strategy": run["allocation_strategy"],
        "total": run["total"],
        "status": run["status"],
        "stopping": is_stopping(RunStatus(run["status"]), stopped_at),
        "stop_requested_at": stopped_at.isoformat() if stopped_at else None,
        "stopped_by": str(run["stopped_by"]) if run["stopped_by"] else None,
        "stopped_by_name": run["stopped_by_name"],
        "started_at": run["started_at"].isoformat(),
        "finished_at": run["finished_at"].isoformat() if run["finished_at"] else None,
        "error": run["error"],
        "started_by": str(started_by) if started_by else None,
        "started_by_name": run["started_by_name"],
        "started_by_avatar_url": run["started_by_avatar_url"],
        # Which lines this run was allowed to dial from. Recorded by
        # `run_numbers_repo.attach` since ADR-8 and never read back until now -
        # a run spread across several numbers is exactly when "where did these
        # calls come from?" is worth being able to answer after the fact, and
        # the per-outcome `from_number_masked` only answers it one call at a
        # time.
        "numbers": numbers,
        "outcomes": outcomes,
        "stats": {
            "completed": len(resolved),
            "total": run["total"],
            "escalated": escalated,
            "auto_closed": sum(1 for o in outcomes if o["disposition"] == "auto_closed"),
            "needs_human_pct": round(100 * escalated / len(outcomes)) if outcomes else 0,
        },
    }


class RunStoppedOut(BaseModel):
    """What the interface needs the moment Stop is pressed.

    Deliberately not the whole run: the detail page is already polling and will
    have the full shape within its next tick. This is the acknowledgement, and
    it carries `stopping` so the button can change immediately rather than
    waiting up to 2.5 seconds to admit it did anything.
    """

    id: str
    status: str
    stopping: bool
    stop_requested_at: str | None
    #: Contacts the dialler had not reached when the stop landed - the people
    #: this actually spares. Excludes anyone mid-conversation, whose call was
    #: already placed and is deliberately left to finish. Approximate by nature:
    #: a call being originated at this instant still connects.
    not_yet_dialled: int


@router.post("/{run_id}/stop", response_model=RunStoppedOut)
async def stop_run(
    run_id: str,
    user: Annotated[CurrentUser, Depends(RequirePermission(Permission.RUNS_START))],
) -> RunStoppedOut:
    """Stop dialling. Calls already in conversation are left to finish.

    **The permission is `runs:start`, not a `runs:stop` of its own.** The set of
    people who may halt a run is exactly the set who may begin one, and a second
    enum member with an identical grant set is a name to keep in step for no
    added control. Stopping is also the strictly less consequential direction -
    it can only ever reduce what an organisation spends and who gets dialled -
    so gating it more tightly than starting would leave the person who began a
    run unable to end it. Viewer holds neither and can do neither.

    **Live calls are not hung up, and that is the product decision.** A stop
    means "dial nobody else"; someone mid-sentence with an agent is not
    disconnected to satisfy a button. Their call ends naturally, its result is
    reported and triaged like any other, and the run closes as `stopped` once
    the last one lands. Everyone not yet reached gets a settled row saying the
    run was stopped before they were dialled, so the count is honest and the run
    can actually close (`RunDialer._stopped_outcome`).

    Idempotent (CLAUDE.md non-negotiable #6). Answers 409 rather than 200 for a
    run that already ended - reporting "Run stopped" for something that stopped
    itself ten minutes ago is a success state for an action that did not happen.
    """
    try:
        run = await request_stop(
            org_id=user.org_id,
            auth_user_id=user.auth_user_id,
            user_id=user.id,
            run_id=run_id,
        )
    except RunNotStoppable as exc:
        # 404 for "no such run" so a caller cannot probe which ids exist; 409 for
        # a real run in the wrong state, which is a genuine conflict and worth
        # telling them about.
        missing = "no longer exists" in str(exc)
        raise HTTPException(
            status_code=status.HTTP_404_NOT_FOUND if missing else status.HTTP_409_CONFLICT,
            detail=str(exc),
        ) from exc

    async with database.as_user(user.auth_user_id) as conn:
        reached = await runs_repo.count_reached(conn, run_id)

    stopped_at = run["stop_requested_at"]
    return RunStoppedOut(
        id=str(run["id"]),
        status=str(run["status"]),
        stopping=is_stopping(RunStatus(str(run["status"])), stopped_at),
        stop_requested_at=stopped_at.isoformat() if stopped_at else None,
        not_yet_dialled=max(0, int(run["total"]) - reached),
    )


#: Refused before a byte is parsed. A workbook this large is either not a
#: contact list or is one that belongs in several uploads - and `openpyxl`
#: materialises far more than the file size in memory.
MAX_UPLOAD_BYTES = 5 * 1024 * 1024

_XLSX_TYPES = frozenset(
    {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel",
        "application/octet-stream",
    }
)


class SheetRowOut(BaseModel):
    """One parsed row - valid, or carrying the reason it is not."""

    row: int
    name: str
    #: Masked. The composer shows enough to recognise a row; the full number
    #: travels only in the `contact` payload the browser posts straight back.
    phone_masked: str
    note: str
    context: dict[str, str]
    valid: bool
    error: str | None
    #: The body `POST /api/v1/runs` accepts for this row, or null when the row
    #: cannot be dialled. Built here so the browser is not re-deriving the
    #: context rules a second time and drifting from them.
    contact: dict[str, Any] | None


class SheetOut(BaseModel):
    rows: list[SheetRowOut]
    #: Context columns found across the sheet, so the composer can say what each
    #: contact brings into its own conversation.
    context_columns: list[str]


@router.post("/parse-sheet", response_model=SheetOut)
async def parse_sheet(
    user: Annotated[CurrentUser, Depends(RequirePermission(Permission.RUNS_START))],
    file: Annotated[UploadFile, File()],
) -> SheetOut:
    """Read an uploaded `.xlsx` into contact rows. Nothing is dialled or stored.

    Parsing happens here rather than in the browser because a workbook is a zip
    of XML with shared strings, styles and typed cells - a phone number typed as
    digits arrives as a float, and getting that wrong turns a valid number into
    `9.19876543210e+11`. CSV is still parsed in the browser, where it is a
    string split and the whole round trip is unnecessary.

    The response is reviewed in the composer before a run is started, exactly as
    a pasted CSV is: an upload is never a dial.
    """
    raw = await file.read()
    if not raw:
        raise HTTPException(status_code=400, detail="That file is empty.")
    if len(raw) > MAX_UPLOAD_BYTES:
        raise HTTPException(
            status_code=413,
            detail=(
                f"That file is larger than {MAX_UPLOAD_BYTES // (1024 * 1024)}MB. "
                "Split it into smaller uploads."
            ),
        )

    filename = (file.filename or "").lower()
    if not filename.endswith((".xlsx", ".xlsm")) and file.content_type not in _XLSX_TYPES:
        raise HTTPException(
            status_code=400,
            detail="Upload an .xlsx file, or paste the rows in as CSV.",
        )

    try:
        rows = _read_workbook(raw)
    except SheetTooLarge as exc:
        raise HTTPException(status_code=413, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    try:
        parsed = parse_rows(rows)
    except SheetTooLarge as exc:
        raise HTTPException(status_code=413, detail=str(exc)) from exc
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc

    columns = sorted({key for row in parsed for key in row.context})
    # Row count and column names only - never a cell value, which is a
    # customer's own data (CLAUDE.md non-negotiable #5).
    log.info(
        "org %s parsed a sheet: %d rows, %d valid, context columns: %s",
        user.org_id,
        len(parsed),
        sum(1 for r in parsed if r.valid),
        ", ".join(columns) or "none",
    )

    return SheetOut(
        rows=[
            SheetRowOut(
                row=row.row,
                name=row.name,
                phone_masked=mask(row.phone) if row.phone else "",
                note=row.note,
                context=row.context,
                valid=row.valid,
                error=row.error,
                contact=to_contact_payload(row) if row.valid else None,
            )
            for row in parsed
        ],
        context_columns=columns,
    )


def _read_workbook(raw: bytes) -> list[list[Any]]:
    """The first worksheet as rows of cell values.

    `read_only` streams rather than building the whole object graph, and
    `data_only` takes a formula cell's cached result - a sheet where the phone
    column is `=CONCAT(...)` would otherwise arrive as the formula text.
    """
    import io

    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except InvalidFileException as exc:
        raise ValueError(
            "That file could not be read as a spreadsheet. Save it as .xlsx and try again."
        ) from exc
    except Exception as exc:
        # openpyxl raises a wide range of types on a corrupt archive, and the
        # vendor's own message is not something to show a user.
        log.warning("could not read an uploaded workbook: %s", type(exc).__name__)
        raise ValueError(
            "That file could not be read as a spreadsheet. Save it as .xlsx and try again."
        ) from exc

    try:
        sheet = workbook.worksheets[0] if workbook.worksheets else None
        if sheet is None:
            raise ValueError("That workbook has no sheets.")

        rows: list[list[Any]] = []
        for row in sheet.iter_rows(values_only=True):
            rows.append(list(row))
            if len(rows) > MAX_ROWS:
                raise SheetTooLarge(
                    f"This sheet has more than {MAX_ROWS:,} rows. "
                    "Split it into smaller uploads."
                )
        return rows
    finally:
        workbook.close()
